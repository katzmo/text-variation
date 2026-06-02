#!/usr/bin/env python3
"""
make_word_guide.py  —  generate alignment_lab_guide.docx
Run once:  python make_word_guide.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── palette ──────────────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1A, 0x44, 0x72)   # headings
MID_BLUE   = RGBColor(0x2E, 0x6D, 0xA8)   # sub-headings
ACCENT     = RGBColor(0xD0, 0x58, 0x00)   # callout labels / highlights
BODY_GREY  = RGBColor(0x33, 0x33, 0x33)
CODE_BG    = RGBColor(0xF2, 0xF2, 0xF2)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_BG    = RGBColor(0xE8, 0xF4, 0xF8)

# ── helpers ───────────────────────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2.2)
    section.bottom_margin = Cm(2.2)

def _to_hex(rgb) -> str:
    """Convert RGBColor to 6-char uppercase hex string. str(RGBColor) returns hex."""
    return str(rgb).upper()

def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _to_hex(rgb))
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        val = kwargs.get(side, {"sz": "4", "val": "single", "color": "CCCCCC"})
        el = OxmlElement(f"w:{side}")
        for k, v in val.items():
            el.set(qn(f"w:{k}"), v)
        tcBorders.append(el)
    tcPr.append(tcBorders)

def h1(text):
    p = doc.add_paragraph()
    p.clear()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = WHITE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Cm(0.4)
    # blue background bar
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1A4472")
    pPr.append(shd)
    return p

def h2(text):
    p = doc.add_paragraph()
    p.clear()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(3)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E6DA8")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def h3(text):
    p = doc.add_paragraph()
    p.clear()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = MID_BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    return p

def body(text, bold_parts=None, indent=0):
    """
    bold_parts: list of substrings to make bold inside text
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if bold_parts:
        remaining = text
        for bp in bold_parts:
            idx = remaining.find(bp)
            if idx == -1:
                continue
            if idx > 0:
                run = p.add_run(remaining[:idx])
                run.font.size = Pt(10.5)
                run.font.color.rgb = BODY_GREY
            run = p.add_run(bp)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = BODY_GREY
            remaining = remaining[idx + len(bp):]
        if remaining:
            run = p.add_run(remaining)
            run.font.size = Pt(10.5)
            run.font.color.rgb = BODY_GREY
    else:
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        run.font.color.rgb = BODY_GREY
    return p

def bullet(text, bold_parts=None, level=1):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    if level == 2:
        p.paragraph_format.left_indent = Cm(1.2)
    if bold_parts:
        remaining = text
        for bp in bold_parts:
            idx = remaining.find(bp)
            if idx == -1:
                continue
            if idx > 0:
                r = p.add_run(remaining[:idx])
                r.font.size = Pt(10.5)
                r.font.color.rgb = BODY_GREY
            r = p.add_run(bp)
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = BODY_GREY
            remaining = remaining[idx + len(bp):]
        if remaining:
            r = p.add_run(remaining)
            r.font.size = Pt(10.5)
            r.font.color.rgb = BODY_GREY
    else:
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        r.font.color.rgb = BODY_GREY
    return p

def code_block(text):
    """Single-line or short code block styled as monospaced shaded paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1A, 0x44, 0x72)
    return p

def callout(label, text):
    """Coloured callout box using a 1-cell table."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, TEAL_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(label + "  ")
    r1.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = ACCENT
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = BODY_GREY
    doc.add_paragraph()  # spacer

def add_data_table(headers, rows, col_widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    # header row
    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, DARK_BLUE)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = WHITE
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # data rows
    for ri, row in enumerate(rows):
        tr = tbl.rows[ri + 1]
        bg = RGBColor(0xF7, 0xF9, 0xFC) if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            if ci == 0:
                r = p.add_run(str(val))
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = DARK_BLUE
            else:
                r = p.add_run(str(val))
                r.font.size = Pt(9.5)
                r.font.color.rgb = BODY_GREY
    # column widths
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in tbl.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()  # spacer
    return tbl

def spacer(pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(pts)

def page_break():
    doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# COVER / TITLE
# ═══════════════════════════════════════════════════════════════════════════════
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_title.add_run("Armenian Manuscript Alignment Lab")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = DARK_BLUE
spacer(2)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_sub.add_run("Step-by-Step User Guide")
r.font.size = Pt(14)
r.font.color.rgb = MID_BLUE
spacer(2)

p_sub2 = doc.add_paragraph()
p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_sub2.add_run("How to run the experiments, read the output, and interpret the results")
r.font.size = Pt(10.5)
r.font.color.rgb = BODY_GREY
spacer(18)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — WHAT THE LAB DOES
# ═══════════════════════════════════════════════════════════════════════════════
h1("1.  What the lab does")
spacer(4)
body(
    "The lab compares Armenian manuscript witnesses of the same text — the Chronicle "
    "of Matthew of Edessa — and tries to match corresponding lines across copies. "
    "Each copy (witness) was written at a different time or place, so the wording can "
    "vary, lines can be added or dropped, and the order can shift slightly. The lab "
    "measures how well different methods find these correspondences.",
    bold_parts=["Chronicle of Matthew of Edessa"],
)
body(
    "A reference copy is called the anchor. All other witnesses are aligned to it. "
    "The lab reports, for each witness, how many of its lines found a match in the anchor "
    "and how confident those matches are.",
    bold_parts=["anchor"],
)
spacer(6)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — FOLDER CONTENTS
# ═══════════════════════════════════════════════════════════════════════════════
h1("2.  What is in the folder")
spacer(4)

add_data_table(
    ["File / Folder", "What it does"],
    [
        ["lab.py",              "Main entry point. Run this to start any experiment."],
        ["parsers.py",          "Reads the TEI XML files and extracts lines of text."],
        ["similarity.py",       "Six ways to score how similar two lines are (Jaccard, Dice, etc.)."],
        ["strategies.py",       "Two alignment algorithms: LIS-based (fast) and Needleman-Wunsch (slow)."],
        ["evaluate.py",         "Calculates alignment rate, mean score, and other metrics."],
        ["data/",               "Folder of 29 TEI XML files — one per manuscript witness."],
        ["results_all.txt",     "All experiment output captured in plain text."],
        ["visualize_results.py","Reads results_all.txt and produces the four PNG plots."],
        ["alignment_results/",  "Folder containing every deliverable (plots, guideline, results)."],
        ["alignment_guideline.md","Plain-English interpretation of the results."],
    ],
    col_widths=[5, 11.5],
)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — SETUP
# ═══════════════════════════════════════════════════════════════════════════════
h1("3.  One-time setup")
spacer(4)
body("Open a terminal, go to the lab folder, and install the required libraries:")
code_block("cd  \"path/to/align-lab\"")
code_block("pip install lxml matplotlib pandas seaborn tabulate")
body(
    "You only need to do this once. After that the lab is ready to run.",
    bold_parts=["only need to do this once"],
)
spacer(6)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — THE FIVE COMMANDS
# ═══════════════════════════════════════════════════════════════════════════════
h1("4.  The five commands and what they do")
spacer(4)

h2("4a.  List available configs")
code_block("python lab.py --list")
body(
    "Prints the names of all named experiment configurations and their settings. "
    "Use this to check what options are available before running anything.",
)
spacer(4)

h2("4b.  Run one named config")
code_block("python lab.py --config baseline")
code_block("python lab.py --config dice")
body(
    "Runs a single named config and prints a comparison table plus sample matches. "
    "Replace baseline or dice with any config name from the list.",
    bold_parts=["single named config"],
)
spacer(4)

h2("4c.  Run all configs at once")
code_block("python lab.py")
body(
    "Runs every config in one go and prints a combined table. "
    "Note: the nw (Needleman-Wunsch) config is extremely slow on large witnesses "
    "(hours). You can comment it out in CONFIGS inside lab.py to skip it.",
    bold_parts=["nw (Needleman-Wunsch) config is extremely slow"],
)
spacer(4)

h2("4d.  All-pairs matrix")
code_block("python lab.py --all-pairs --config baseline")
code_block("python lab.py --all-pairs --config dice")
body(
    "Aligns every witness against every other witness and prints a grid. "
    "Each cell shows the alignment rate (%) when the row witness is aligned "
    "to the column anchor. This takes about 7 minutes per run because there "
    "are 29 witnesses x 28 anchors = 812 alignment jobs.",
    bold_parts=["every witness against every other witness", "7 minutes"],
)
spacer(4)

h2("4e.  Visualise the results")
code_block("python visualize_results.py")
body(
    "Reads results_all.txt and produces the four PNG plots and a summary table. "
    "Run this after all experiments are complete.",
    bold_parts=["Run this after all experiments are complete"],
)
spacer(6)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — CONFIG REFERENCE
# ═══════════════════════════════════════════════════════════════════════════════
h1("5.  Config reference table")
spacer(4)
body("Each named config combines a similarity measure, a section window size, and an alignment strategy.")
spacer(4)

add_data_table(
    ["Config", "Similarity", "Window", "Strategy", "What it tests"],
    [
        ["baseline",    "Jaccard",    "1 line",    "LIS",  "Current production default."],
        ["dice",        "Dice",       "1 line",    "LIS",  "Rewards shared words more. Best overall."],
        ["levenshtein", "Levenshtein","1 line",    "LIS",  "Word-order aware. Selective."],
        ["combined",    "50% Jac + 50% Lev","1 line","LIS","Balanced blend. Safe middle ground."],
        ["sections5",   "Jaccard",    "5 lines",   "LIS",  "Merges 5 lines before comparing."],
        ["sections10",  "Jaccard",    "10 lines",  "LIS",  "Merges 10 lines. Coarse-grained."],
        ["positional",  "Jaccard",    "1 line",    "LIS",  "Soft penalty for far-away matches."],
        ["banded",      "Jaccard",    "1 line",    "LIS",  "Hard positional band. Cuts false positives."],
        ["nw",          "Jaccard",    "1 line",    "NW",   "Global alignment. SLOW on large texts."],
    ],
    col_widths=[3, 3.8, 2.2, 2, 5.5],
)

callout(
    "Tip:",
    "Jaccard measures which words the two lines share, ignoring word order. "
    "Dice does the same but counts shared words more generously. "
    "Levenshtein checks whether the words appear in the same sequence."
)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 6 — READING THE OUTPUT
# ═══════════════════════════════════════════════════════════════════════════════
h1("6.  How to read the output")
spacer(4)

h2("6a.  The comparison table")
body(
    "When you run python lab.py --config baseline, you see something like this:",
)
code_block("Comparing configs   witness=M1896   anchor=M1767")
code_block("")
code_block("config    rate    aligned  mean_sc  mono  anc_cov  time")
code_block("--------  ------  -------  -------  ----  -------  -----")
code_block("baseline  18.9%   3374     0.505    1.00  9.9%     0.46s")
spacer(4)

add_data_table(
    ["Column", "Meaning"],
    [
        ["rate",    "Alignment rate. The percentage of witness lines that found a match in the anchor. Higher = more lines matched."],
        ["aligned", "The raw count of lines that were matched (numerator of rate)."],
        ["mean_sc", "Mean similarity score of the matched pairs (0 to 1). Higher = matches are more confident."],
        ["mono",    "Monotonicity. 1.00 means all matched lines appear in the same reading order as the anchor. Near 1.00 is normal and good."],
        ["anc_cov", "Anchor coverage. The percentage of anchor lines that were matched by at least one witness line."],
        ["time",    "Computation time for this config only (not counting corpus loading)."],
    ],
    col_widths=[2.8, 13.7],
)

callout(
    "Important:",
    "A higher alignment rate is not automatically better. It could also mean more "
    "false positives (lines matched incorrectly). Always look at the sample matches "
    "at the bottom of the output to judge quality."
)

h2("6b.  The agreement table")
body(
    "When you run all configs together (python lab.py), you also see an Agreement "
    "table like this:",
)
code_block("Agreement with baseline (same target where both align):")
code_block("config        agree")
code_block("----------    -----")
code_block("dice          94%")
spacer(2)
body(
    "This tells you: when both dice and baseline matched the same witness line, "
    "did they agree on which anchor line it maps to? "
    "94% means they almost always picked the same target. "
    "High agreement means both methods are finding real correspondences.",
    bold_parts=["High agreement means both methods are finding real correspondences."],
)
spacer(4)

h2("6c.  Sample matches")
body(
    "Below each table you see up to 8 matched pairs, highest score first:",
)
code_block("1.000  կու ի վր երկրի          |  կու ի վր երկրի")
code_block("0.857  նր և ի գիշերին ընդ մի  |  նր և ի գիշերին ընդ մի")
spacer(2)
body(
    "The left column is the witness line. The right column is the anchor line. "
    "A score of 1.000 means the lines are identical after cleaning. "
    "Scores above 0.7 are strong matches. Scores near the threshold (0.35) need "
    "manual verification.",
    bold_parts=["A score of 1.000 means the lines are identical after cleaning.",
                "Scores above 0.7 are strong matches.",
                "Scores near the threshold (0.35) need manual verification."],
)
spacer(6)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 7 — ALL-PAIRS MATRIX
# ═══════════════════════════════════════════════════════════════════════════════
h1("7.  Reading the all-pairs matrix")
spacer(4)
body(
    "The matrix has witnesses as rows and anchors as columns. Each cell is the "
    "alignment rate (%) when that row-witness is aligned to that column-anchor.",
    bold_parts=["row-witness", "column-anchor"],
)
code_block("wit \\ anc   BZ430  BZ449  BZ644  ...  M3380  ...")
code_block("---------   -----  -----  -----  ...  -----  ...")
code_block("BZ430          --    15%    14%  ...    12%  ...")
code_block("BZ449           1%   --     48%  ...    36%  ...")
spacer(4)

h2("What to look for")
bullet(
    "Dense column = good anchor.  A column with many high percentages means "
    "many witnesses align well to that anchor. This is a sign the anchor covers "
    "the text broadly and uses representative vocabulary.",
    bold_parts=["Dense column = good anchor."],
)
bullet(
    "Sparse column = bad anchor.  If most cells in a column are below 10%, that "
    "witness is a poor reference. It may cover a different section, or use very "
    "different vocabulary.",
    bold_parts=["Sparse column = bad anchor."],
)
bullet(
    "Very high mutual rate = related copies.  If witness A aligns to witness B "
    "at 70-85% AND B aligns to A at similar rates, they may be copies of the same "
    "source manuscript.",
    bold_parts=["Very high mutual rate = related copies."],
)
bullet(
    "Low row = fragment witness.  If a witness row has low values everywhere "
    "(under 15%), that manuscript is likely a fragment covering only part of the text.",
    bold_parts=["Low row = fragment witness."],
)
spacer(4)

callout(
    "Key finding:",
    "In this corpus, M3380 (12,425 lines) is the best anchor — it gets the "
    "highest mean alignment rate across all other witnesses (38% baseline, 57% dice). "
    "M1767 is the longest witness (33,984 lines) but the worst anchor "
    "(only 7-9% mean rate), meaning it covers a different part of the text."
)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 8 — THE FOUR PLOTS
# ═══════════════════════════════════════════════════════════════════════════════
h1("8.  The four plots explained")
spacer(4)

h2("plot_config_comparison.png  — all configs side by side")
body(
    "Two grouped bars per config: alignment rate (blue) and mean score scaled x100 "
    "(green). Use this to compare methods at a glance. Look for configs where both "
    "bars are tall: high rate AND high confidence.",
    bold_parts=["high rate AND high confidence"],
)
spacer(4)

h2("plot_similarity_measures.png  — zoom in on the four similarity measures")
body(
    "Shows only baseline (Jaccard), Dice, Levenshtein, and Combined. "
    "A focused view for answering: which word-comparison method works best?",
)
spacer(4)

h2("plot_section_size.png  — effect of merging lines into sections")
body(
    "A line chart with window size on the x-axis (1, 5, 10 lines) and both metrics "
    "on the y-axes. A downward slope on both lines means larger sections hurt alignment.",
    bold_parts=["downward slope on both lines means larger sections hurt alignment"],
)
body(
    "Note: the y-axis rates are not directly comparable across window sizes. "
    "A 12.7% rate with window=5 means 12.7% of 5-line sections matched, not "
    "12.7% of individual lines.",
    bold_parts=["Note:"],
)
spacer(4)

h2("plot_allpairs_heatmap.png  — 29 x 29 witness matrix")
body(
    "A colour grid where dark orange or red = high alignment rate. "
    "Look for dark columns (good anchors) and pale rows (fragment witnesses). "
    "Bright diagonal cells would mean a witness aligns to itself — these are "
    "shown as blank because self-alignment is not measured.",
    bold_parts=["dark columns (good anchors)", "pale rows (fragment witnesses)"],
)
spacer(6)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 9 — STEP-BY-STEP ANALYSIS WORKFLOW
# ═══════════════════════════════════════════════════════════════════════════════
h1("9.  Step-by-step analysis workflow")
spacer(4)
body("Follow these steps to go from raw files to a complete analysis.")
spacer(4)

h2("Step 1 — run the individual configs")
body("Run each config separately and append all output to one file:")
code_block("python lab.py --config baseline   >> results_all.txt 2>&1")
code_block("python lab.py --config dice        >> results_all.txt 2>&1")
code_block("python lab.py --config levenshtein >> results_all.txt 2>&1")
code_block("python lab.py --config combined    >> results_all.txt 2>&1")
code_block("python lab.py --config sections5   >> results_all.txt 2>&1")
code_block("python lab.py --config sections10  >> results_all.txt 2>&1")
code_block("python lab.py --config positional  >> results_all.txt 2>&1")
code_block("python lab.py --config banded      >> results_all.txt 2>&1")
body(
    "Each command takes about 30-60 seconds (most of the time is loading the 29 XML files).",
)
spacer(4)

h2("Step 2 — run the all-pairs matrix (optional, takes ~15 minutes)")
code_block("python lab.py --all-pairs --config baseline >> results_all.txt 2>&1")
code_block("python lab.py --all-pairs --config dice     >> results_all.txt 2>&1")
body("This produces the witness-vs-witness grids needed for the heatmap plot.")
spacer(4)

h2("Step 3 — produce the plots and summary table")
code_block("python visualize_results.py")
body(
    "This reads results_all.txt and writes four PNG files to the current folder. "
    "It also prints a summary table to the terminal.",
)
spacer(4)

h2("Step 4 — copy everything into the results folder")
code_block("mkdir -p alignment_results")
code_block("cp results_all.txt alignment_guideline.md visualize_results.py plot_*.png alignment_results/")
spacer(6)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 10 — INTERPRETING THE NUMBERS
# ═══════════════════════════════════════════════════════════════════════════════
h1("10.  How to interpret the numbers")
spacer(4)

h2("Alignment rate: what counts as good?")
body(
    "There is no fixed target. For the largest witness pair tested "
    "(M1896 vs M1767), even the best config (dice) only matched 25.7% of lines. "
    "This is expected: the two witnesses may cover different chapters.",
)
add_data_table(
    ["Rate range", "Interpretation"],
    [
        ["70 – 100%", "The two witnesses cover the same section of the text and use similar wording. Expected for closely related copies (e.g. M2644 and M2855 in this corpus)."],
        ["30 – 70%",  "Substantial overlap. The witnesses share much of the text but differ in some sections or wording."],
        ["10 – 30%",  "Partial overlap. Either the witnesses cover different parts of the text, or their vocabulary diverges significantly."],
        ["0 – 10%",   "Very low overlap. One witness may be a fragment, or the two witnesses cover completely different sections."],
    ],
    col_widths=[2.5, 14.0],
)

h2("Mean score: what counts as a confident match?")
add_data_table(
    ["Score range", "Interpretation"],
    [
        ["0.8 – 1.0", "High confidence. The lines are nearly identical or share most words. Trust these matches."],
        ["0.5 – 0.8", "Moderate confidence. The lines share key vocabulary but may differ in some words. Verify a sample."],
        ["0.35 – 0.5","Low confidence (near the threshold). These are the weakest accepted matches. Inspect manually."],
        ["< 0.35",    "Below threshold, not reported as a match."],
    ],
    col_widths=[2.5, 14.0],
)
spacer(4)

callout(
    "Rule of thumb:",
    "If mean_sc is above 0.55 AND rate is above 20%, the config is performing well "
    "on this witness pair. Dice meets both criteria on M1896 vs M1767 "
    "(rate 25.7%, mean_sc 0.612)."
)
spacer(4)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 11 — RESULTS FROM THIS CORPUS
# ═══════════════════════════════════════════════════════════════════════════════
h1("11.  Results from this corpus (M1896 vs M1767)")
spacer(4)
body(
    "All single-anchor tests used M1896 (17,866 lines) as the witness and "
    "M1767 (33,984 lines) as the anchor. These are the two longest witnesses."
)
spacer(4)

add_data_table(
    ["Config", "Align rate", "Mean score", "Verdict"],
    [
        ["baseline (Jaccard)", "18.9%", "0.505", "Starting point. Acceptable."],
        ["dice",               "25.7%", "0.612", "Best overall. +36% more matches, higher confidence."],
        ["levenshtein",        "20.5%", "0.510", "Slight gain over baseline. Slower. Adds no clear benefit here."],
        ["combined",           "21.2%", "0.495", "Marginal improvement in rate but lower confidence than dice."],
        ["banded",             "23.5%", "0.499", "Hard positional band helps. Good second choice."],
        ["positional",         "20.7%", "0.495", "Soft prior adds little benefit on this pair."],
        ["sections5",          "12.7%", "0.405", "Section-level rate. Fewer and weaker matches than line-level."],
        ["sections10",          "5.2%", "0.381", "Very low. Coarse sections hurt badly on this text."],
    ],
    col_widths=[4.5, 2.5, 2.8, 6.7],
)

h2("All-pairs: best and worst anchors (baseline config)")
add_data_table(
    ["Rank", "Witness", "Lines", "Mean col rate", "Note"],
    [
        ["1st (best)", "M3380", "12,425", "38.4%", "Best overall anchor. Dense column."],
        ["2nd",        "M3520", "11,702", "36.7%", "Very close second. Either works well."],
        ["3rd",        "OX E",  "10,410", "33.8%", "Strong anchor, also well-connected."],
        ["…",          "…",     "…",      "…",     "…"],
        ["Worst",      "M1767", "33,984",  "7.5%", "Longest witness, worst anchor. Avoid."],
    ],
    col_widths=[2.5, 2.5, 2.5, 3.5, 5.5],
)
spacer(4)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 12 — RECOMMENDED PRODUCTION SETTINGS
# ═══════════════════════════════════════════════════════════════════════════════
h1("12.  Recommended production settings")
spacer(4)

body(
    "Based on all experiments, these are the settings that should give the best "
    "results for aligning this corpus:",
    bold_parts=["best results"],
)
spacer(4)

add_data_table(
    ["Setting", "Value", "Why"],
    [
        ["Similarity measure", "dice",     "Finds 36% more matches than Jaccard with higher mean score."],
        ["Window size",        "1 (lines)","Line-level alignment outperforms section-level on this text."],
        ["Strategy",           "lis",      "Fast, order-preserving. Needleman-Wunsch is too slow on large witnesses."],
        ["Positional band",    "band_frac=0.15","Hard band removes false positives from formulaic phrases."],
        ["Anchor witness",     "M3380",    "Highest mean alignment rate across all other witnesses (38-57%)."],
        ["Threshold",          "0.35",     "Default. Works well across all configs. Do not lower without checking false positives."],
    ],
    col_widths=[4, 3.5, 9.0],
)

body("To add this as a named config, open lab.py and add one line in the CONFIGS dictionary:")
code_block('\"dice_banded\": dict(window=1, similarity=\"dice\", strategy=\"lis\", band_frac=0.15),')
body("Then run it with:")
code_block("python lab.py --config dice_banded")
spacer(6)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 13 — KNOWN LIMITATIONS
# ═══════════════════════════════════════════════════════════════════════════════
h1("13.  Known limitations")
spacer(4)

bullet(
    "Low rates are correct for fragments.  Witnesses BZ430 (496 lines) and M6686 "
    "(798 lines) cover only a small part of the chronicle. Their alignment rates "
    "will always be low. This is a correct result, not a failure.",
    bold_parts=["Low rates are correct for fragments."],
)
bullet(
    "Monster lines in some witnesses.  A small number of lines in M1896 "
    "and M1767 contain 15,000-56,000 words due to a parser edge case (the last "
    "<lb/> milestone element capturing all remaining text). A safeguard has been "
    "added to similarity.py: the Levenshtein function falls back to Jaccard for "
    "lines longer than 300 words. Without this fix, those lines would cause "
    "the alignment to hang for hours.",
    bold_parts=["Monster lines in some witnesses.", "safeguard has been added"],
)
bullet(
    "Needleman-Wunsch (nw) does not scale.  For witnesses with 10,000+ lines, "
    "the NW algorithm needs several gigabytes of memory and many hours to complete. "
    "Use it only on small witness pairs (under about 1,000 lines each).",
    bold_parts=["Needleman-Wunsch (nw) does not scale."],
)
bullet(
    "Section rates are not comparable to line rates.  A 12.7% alignment rate "
    "with window=5 means 12.7% of 5-line sections matched, not 12.7% of individual "
    "lines. Do not compare these numbers directly.",
    bold_parts=["Section rates are not comparable to line rates."],
)
bullet(
    "No gold standard.  Without a manually verified reference alignment, "
    "a higher rate could also mean more false positives. Always inspect the sample "
    "matches to judge quality. The agreement column (cross-config comparison) is "
    "a useful proxy.",
    bold_parts=["No gold standard.", "Always inspect the sample matches"],
)
bullet(
    "Anchor-free alignment not available.  The --mode columns and "
    "--mode progressive commands require a multi.py file that has not yet been "
    "implemented.",
    bold_parts=["Anchor-free alignment not available."],
)
spacer(6)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 14 — QUICK-REFERENCE CHEAT SHEET
# ═══════════════════════════════════════════════════════════════════════════════
h1("14.  Quick-reference cheat sheet")
spacer(4)

add_data_table(
    ["Task", "Command"],
    [
        ["See all config names",                 "python lab.py --list"],
        ["Run one config",                       "python lab.py --config dice"],
        ["Run all configs",                      "python lab.py"],
        ["All-pairs matrix (baseline)",          "python lab.py --all-pairs --config baseline"],
        ["All-pairs matrix (dice)",              "python lab.py --all-pairs --config dice"],
        ["Produce plots + summary",              "python visualize_results.py"],
        ["Use a different data folder",          "python lab.py --data path/to/folder"],
        ["Control number of sample matches",     "python lab.py --config baseline --samples 12"],
        ["Save all output to a file",            "python lab.py --config dice >> results.txt 2>&1"],
    ],
    col_widths=[7, 9.5],
)
spacer(4)

callout(
    "Where to find the results:",
    "All deliverables are in the alignment_results/ folder inside align-lab/. "
    "Open alignment_guideline.md for a short plain-language interpretation."
)

# ── Save ──────────────────────────────────────────────────────────────────────
output_path = "alignment_results/alignment_lab_guide.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
